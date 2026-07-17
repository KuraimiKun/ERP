from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE = Path(r"f:\erpdjango\thesis")
OUT = BASE / "sparsity_thesis_polished_final.docx"

MAIN_TEX = BASE / "sparsity_thesis.tex"
CHAPTERS = [
    BASE / "chapters" / "Chapter_1_Introduction_v2.tex",
    BASE / "chapters" / "Chapter_2_LiteratureReview_v2.tex",
    BASE / "chapters" / "Chapter_3_Methodology_v2.tex",
    BASE / "chapters" / "Chapter_4_Results_v2.tex",
    BASE / "chapters" / "Chapter_5_Conclusion_v2.tex",
]

GRAPHIC_DIRS = [
    BASE / "figures",
    BASE / "results" / "sparsity_experiment" / "figures",
    BASE / "results" / "hr_attrition_experiment" / "figures",
    BASE / "results" / "lead_scoring_experiment" / "figures",
]


def remove_comments(text: str) -> str:
    out = []
    for line in text.splitlines():
        if line.lstrip().startswith("%"):
            continue
        line = re.sub(r"(?<!\\)%.*$", "", line)
        out.append(line)
    return "\n".join(out)


def resolve_figure_path(raw: str) -> Path | None:
    raw = raw.strip()
    candidates = [raw]
    if "." not in Path(raw).name:
        candidates.extend([f"{raw}.png", f"{raw}.jpg", f"{raw}.jpeg", f"{raw}.webp", f"{raw}.pdf"])
    for name in candidates:
        p = Path(name)
        if p.is_absolute() and p.exists():
            return p
        for d in GRAPHIC_DIRS:
            full = d / p.name if len(p.parts) == 1 else BASE / p
            if full.exists() and full.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
                return full
    return None


def tex_inline_cleanup(text: str) -> str:
    # Remove control commands that should never appear in prose output.
    text = re.sub(r"\\addcontentsline\{[^{}]*\}\{[^{}]*\}\{[^{}]*\}", "", text)
    text = re.sub(r"\\markboth\{[^{}]*\}\{[^{}]*\}", "", text)
    text = re.sub(r"\\tableofcontents|\\listoffigures|\\listoftables", "", text)

    text = re.sub(r"\\cite[t|p]?\*?\{([^{}]*)\}", r"[\1]", text)
    text = re.sub(r"\\ref\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\label\{[^{}]*\}", "", text)
    replacements = [
        (r"\\textbf\{([^{}]*)\}", r"\1"),
        (r"\\textit\{([^{}]*)\}", r"\1"),
        (r"\\emph\{([^{}]*)\}", r"\1"),
        (r"\\texttt\{([^{}]*)\}", r"\1"),
        (r"\\mathbf\{([^{}]*)\}", r"\1"),
        (r"\\mathsf\{([^{}]*)\}", r"\1"),
        (r"\\mathrm\{([^{}]*)\}", r"\1"),
        (r"\\mathtt\{([^{}]*)\}", r"\1"),
    ]
    for pattern, repl in replacements:
        text = re.sub(pattern, repl, text)

    # Common inline math symbols used in the thesis.
    text = text.replace(r"\rho", "rho")
    text = text.replace(r"\tau", "tau")
    text = text.replace(r"\sigma", "sigma")

    # Normalize grouped number formatting like 50{,}000 -> 50,000.
    text = re.sub(r"(\d)\{,\}(\d)", r"\1,\2", text)

    # Remove leftover simple command forms while keeping content where possible.
    text = re.sub(r"\\[a-zA-Z]+\*?\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[a-zA-Z]+", "", text)

    # Normalize escaped punctuation and leftover TeX separators.
    text = text.replace(r"\\", " ")
    text = text.replace(r"\"", '"')
    text = text.replace(r"\;", " ")
    text = text.replace(r"\,", " ")

    # Remove any final residual backslash escapes.
    text = re.sub(r"\\([A-Za-z0-9_]+)", r"\1", text)
    text = text.replace("\\", "")

    # Drop braces that are formatting-only after command removal.
    text = text.replace("{", "").replace("}", "")

    text = text.replace(r"\_", "_")
    text = text.replace(r"\&", "&")
    text = text.replace(r"\%", "%")
    text = text.replace("~", " ")
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    text = text.replace("``", '"').replace("''", '"')
    return re.sub(r"\s+", " ", text).strip()


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update field to populate Table of Contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(placeholder)
    run._r.append(fld_end)


def parse_title_and_abstract() -> tuple[str, str]:
    src = remove_comments(MAIN_TEX.read_text(encoding="utf-8"))
    title_match = re.search(r"\\newcommand\{\\projecttitle\}\{(.+?)\}", src, flags=re.S)
    title = tex_inline_cleanup(title_match.group(1) if title_match else "Thesis")
    title = title.replace("\\\\", ": ").replace("\\", " ")
    abs_match = re.search(r"\\chapter\*\{Abstract\}(.*?)\\noindent\\textbf\{Keywords:\}", src, flags=re.S)
    abstract = tex_inline_cleanup(abs_match.group(1) if abs_match else "")
    return title, abstract


def tokenize_blocks(src: str) -> list[tuple[str, str]]:
    src = remove_comments(src)
    lines = src.splitlines()
    blocks: list[tuple[str, str]] = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        m = re.match(r"\\chapter\*?\{(.+?)\}", line)
        if m:
            blocks.append(("h1", tex_inline_cleanup(m.group(1))))
            i += 1
            continue
        m = re.match(r"\\section\*?\{(.+?)\}", line)
        if m:
            blocks.append(("h2", tex_inline_cleanup(m.group(1))))
            i += 1
            continue
        m = re.match(r"\\subsection\*?\{(.+?)\}", line)
        if m:
            blocks.append(("h3", tex_inline_cleanup(m.group(1))))
            i += 1
            continue
        m = re.match(r"\\subsubsection\*?\{(.+?)\}", line)
        if m:
            blocks.append(("h4", tex_inline_cleanup(m.group(1))))
            i += 1
            continue
        m = re.match(r"\\paragraph\*?\{(.+?)\}", line)
        if m:
            blocks.append(("lead", tex_inline_cleanup(m.group(1))))
            i += 1
            continue

        if line.startswith("\\item"):
            item = tex_inline_cleanup(re.sub(r"^\\item(?:\[.*?\])?", "", line).strip())
            blocks.append(("li", item))
            i += 1
            continue

        if line.startswith("\\begin{equation") or line.startswith("\\begin{align"):
            env = "equation" if "equation" in line else "align"
            buf = []
            i += 1
            while i < len(lines) and f"\\end{{{env}" not in lines[i]:
                buf.append(lines[i])
                i += 1
            blocks.append(("eq", tex_inline_cleanup(" ".join(buf))))
            i += 1
            continue

        if line.startswith("\\begin{lstlisting"):
            buf = []
            i += 1
            while i < len(lines) and "\\end{lstlisting}" not in lines[i]:
                buf.append(lines[i].rstrip("\n"))
                i += 1
            blocks.append(("code", "\n".join(buf).strip()))
            i += 1
            continue

        if "\\includegraphics" in line:
            joined = line
            j = i + 1
            while j < len(lines) and "}" not in joined:
                joined += lines[j]
                j += 1
            g = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", joined)
            if g:
                blocks.append(("fig", g.group(1).strip()))
            i = j
            continue

        if line.startswith("\\begin{figure"):
            fig_buf = [line]
            i += 1
            while i < len(lines) and "\\end{figure}" not in lines[i]:
                fig_buf.append(lines[i])
                i += 1
            fig_buf.append(lines[i] if i < len(lines) else "")
            blob = "\n".join(fig_buf)
            g = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", blob)
            c = re.search(r"\\caption\{(.+?)\}", blob, flags=re.S)
            if g:
                blocks.append(("fig", g.group(1).strip()))
            if c:
                blocks.append(("caption", tex_inline_cleanup(c.group(1))))
            i += 1
            continue

        if line.startswith("\\begin{") or line.startswith("\\end{"):
            i += 1
            continue

        # Merge paragraph until control boundary.
        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt.startswith("\\") or "\\includegraphics" in nxt:
                break
            para.append(nxt)
            i += 1
        blocks.append(("p", tex_inline_cleanup(" ".join(para))))

    return blocks


def set_base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    for style_name in ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_title_page(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)

    for text, size, bold in [
        ("A Bachelor's Thesis", 14, False),
        ("Ahmed Ameen Hassan Alkuraimi", 14, True),
        ("NIM / 220535618354", 12, False),
        ("Faculty of Engineering", 12, False),
        ("State University of Malang", 12, False),
        ("2026", 12, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(text)
        rr.font.size = Pt(size)
        rr.bold = bold

    logo = resolve_figure_path("Lambang-UM.png")
    if logo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(logo), width=Inches(1.5))

    doc.add_page_break()


def add_abstract_and_toc(doc: Document, abstract: str) -> None:
    doc.add_heading("Abstract", level=1)
    if abstract:
        doc.add_paragraph(abstract)
    doc.add_paragraph(
        "Keywords: Explainable AI, SHAP, LIME, multi-tenant SaaS, ERP, cold-start, data sparsity, Django, reliability cliff."
    )

    doc.add_heading("Table of Contents", level=1)
    add_toc_field(doc)
    doc.add_page_break()


def parse_bib_entries(bib_path: Path) -> list[str]:
    if not bib_path.exists():
        return []
    src = bib_path.read_text(encoding="utf-8")
    entries = re.findall(r"@\w+\s*\{\s*[^,]+\s*,(.*?)\n\}\s*", src, flags=re.S)

    def field(body: str, name: str) -> str:
        m = re.search(rf"{name}\s*=\s*(\{{.*?\}}|\".*?\")\s*,", body, flags=re.S | re.I)
        if not m:
            return ""
        val = m.group(1).strip()
        if (val.startswith("{") and val.endswith("}")) or (val.startswith('"') and val.endswith('"')):
            val = val[1:-1]
        return tex_inline_cleanup(val)

    out: list[str] = []
    for body in entries:
        author = field(body, "author") or "Unknown author"
        year = field(body, "year") or "n.d."
        title = field(body, "title") or "Untitled"
        journal = field(body, "journal") or field(body, "booktitle") or field(body, "publisher")
        author = author.replace(" and ", "; ")
        parts = [f"{author} ({year}).", f"{title}."]
        if journal:
            parts.append(f"{journal}.")
        out.append(" ".join(parts))
    return out


def add_blocks(doc: Document, blocks: Iterable[tuple[str, str]]) -> None:
    for kind, payload in blocks:
        if not payload:
            continue
        if kind == "h1":
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "h3":
            doc.add_heading(payload, level=3)
        elif kind == "h4":
            doc.add_heading(payload, level=4)
        elif kind == "lead":
            p = doc.add_paragraph()
            r = p.add_run(payload)
            r.bold = True
        elif kind == "li":
            doc.add_paragraph(payload, style="List Bullet")
        elif kind == "eq":
            p = doc.add_paragraph("Equation: " + payload)
            p.style = doc.styles["Intense Quote"] if "Intense Quote" in doc.styles else doc.styles["Normal"]
        elif kind == "code":
            if payload.strip():
                p = doc.add_paragraph(payload)
                p.style = doc.styles["No Spacing"] if "No Spacing" in doc.styles else doc.styles["Normal"]
        elif kind == "fig":
            path = resolve_figure_path(payload)
            if path:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(path), width=Inches(5.8))
            else:
                doc.add_paragraph(f"[Figure not found: {payload}]")
        elif kind == "caption":
            p = doc.add_paragraph(payload)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].italic = True
        else:
            doc.add_paragraph(payload)


def add_footer_pagenum(doc: Document) -> None:
    for section in doc.sections:
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = fp.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")

        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)


def polish_paragraph_layout(doc: Document) -> None:
    for para in doc.paragraphs:
        style_name = para.style.name if para.style is not None else ""
        text = para.text.strip()

        para.paragraph_format.space_after = Pt(6)

        if style_name.startswith("Heading"):
            para.paragraph_format.space_before = Pt(10)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.first_line_indent = None
            continue

        if style_name.startswith("List") or text.startswith("["):
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.first_line_indent = None
            continue

        if para.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.first_line_indent = None
            continue

        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.first_line_indent = Inches(0.28)


def build_docx() -> Path:
    title, abstract = parse_title_and_abstract()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1.18)
    sec.bottom_margin = Inches(1.18)
    sec.left_margin = Inches(1.57)
    sec.right_margin = Inches(1.18)

    set_base_styles(doc)
    add_title_page(doc, title)
    add_abstract_and_toc(doc, abstract)

    for chapter in CHAPTERS:
        blocks = tokenize_blocks(chapter.read_text(encoding="utf-8"))
        add_blocks(doc, blocks)
        doc.add_page_break()

    doc.add_heading("References", level=1)
    bib_entries = parse_bib_entries(BASE / "references_sparsity.bib")
    if bib_entries:
        for entry in bib_entries:
            p = doc.add_paragraph(entry)
            p.style = "List Number"
    else:
        doc.add_paragraph("References are maintained in references_sparsity.bib.")

    polish_paragraph_layout(doc)

    add_footer_pagenum(doc)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    written = build_docx()
    print(f"WROTE {written}")
