"""
ERP Data Export Service

Provides CSV and Excel export for all list views.
Each entity defines its exportable columns and filter logic.
"""

import csv
import hashlib
from datetime import date, datetime
from io import BytesIO

from django.core.cache import cache
from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from erp.models import (
    Customer, Invoice, Payment, Expense, ExpenseCategory,
    Account, TaxRate, PaymentMethod, Currency, FiscalYear,
    Lead, Campaign, Task,
    Employee, Department, Attendance, LeaveRequest, Payroll,
    Supplier, PurchaseOrder,
    Quote, ServiceCategory, ServicePricing,
    Product, ProductCategory,
    JournalEntry,
)


# ── Column definitions ──────────────────────────────────────────────────────
# Each entry: (header, field_or_callable)
# - str field: dot notation for FK traversal, e.g. 'customer.name'
# - callable: lambda obj: value
# - str starting with 'get_': calls the method, e.g. 'get_status_display'

def _resolve_field(obj, field):
    """Resolve a field spec to a value from a model instance."""
    if callable(field):
        return field(obj)
    parts = field.split('.')
    val = obj
    for part in parts:
        if val is None:
            return ''
        if part.startswith('get_') and part.endswith('_display'):
            val = getattr(val, part, lambda: '')()
        else:
            val = getattr(val, part, '')
    if isinstance(val, bool):
        return 'Yes' if val else 'No'
    if isinstance(val, (date, datetime)):
        return val.isoformat() if val else ''
    if val is None:
        return ''
    return val


# ── Export registry ─────────────────────────────────────────────────────────

EXPORT_REGISTRY = {
    'customers': {
        'model': Customer,
        'filename': 'customers',
        'select_related': ['currency'],
        'columns': [
            ('Customer ID', 'customer_id'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Type', 'get_customer_type_display'),
            ('Email', 'email'),
            ('Phone', 'phone'),
            ('Mobile', 'mobile'),
            ('Tax ID', 'tax_id'),
            ('Credit Limit', 'credit_limit'),
            ('Balance', 'current_balance'),
            ('Payment Terms (days)', 'payment_terms'),
            ('Currency', 'currency.code'),
            ('City', 'billing_city'),
            ('Country', 'billing_country'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: _filter_customers(request, qs),
    },
    'invoices': {
        'model': Invoice,
        'filename': 'invoices',
        'select_related': ['customer', 'currency'],
        'columns': [
            ('Invoice #', 'invoice_number'),
            ('Customer', 'customer.name'),
            ('Type', 'get_invoice_type_display'),
            ('Date', 'invoice_date'),
            ('Due Date', 'due_date'),
            ('Subtotal', 'subtotal'),
            ('Discount', 'discount_amount'),
            ('Tax', 'tax_amount'),
            ('Total', 'total_amount'),
            ('Paid', 'paid_amount'),
            ('Balance Due', 'balance_due'),
            ('Status', 'get_status_display'),
            ('Currency', 'currency.code'),
            ('Reference', 'reference'),
        ],
        'filters': lambda request, qs: _filter_invoices(request, qs),
    },
    'payments': {
        'model': Payment,
        'filename': 'payments',
        'select_related': ['customer', 'payment_method', 'currency'],
        'columns': [
            ('Payment #', 'payment_number'),
            ('Customer', 'customer.name'),
            ('Type', 'get_payment_type_display'),
            ('Date', 'payment_date'),
            ('Method', 'payment_method.name'),
            ('Amount', 'amount'),
            ('Allocated', 'allocated_amount'),
            ('Unallocated', 'unallocated_amount'),
            ('Currency', 'currency.code'),
            ('Reference', 'reference_number'),
        ],
        'filters': lambda request, qs: _filter_payments(request, qs),
    },
    'expenses': {
        'model': Expense,
        'filename': 'expenses',
        'select_related': ['category', 'employee', 'currency'],
        'columns': [
            ('Expense #', 'expense_number'),
            ('Date', 'expense_date'),
            ('Category', 'category.name'),
            ('Employee', lambda o: str(o.employee) if o.employee else ''),
            ('Description', 'description'),
            ('Amount', 'amount'),
            ('Tax', 'tax_amount'),
            ('Total', 'total_amount'),
            ('Status', 'get_status_display'),
            ('Currency', 'currency.code'),
        ],
        'filters': lambda request, qs: _filter_expenses(request, qs),
    },
    'accounts': {
        'model': Account,
        'filename': 'chart_of_accounts',
        'select_related': ['parent'],
        'columns': [
            ('Code', 'code'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Type', 'get_account_type_display'),
            ('Parent', 'parent.code'),
            ('Opening Balance', 'opening_balance'),
            ('Current Balance', 'current_balance'),
            ('System', 'is_system'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: _filter_accounts(request, qs),
    },
    'tax_rates': {
        'model': TaxRate,
        'filename': 'tax_rates',
        'select_related': [],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Rate (%)', 'rate'),
            ('Default', 'is_default'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'payment_methods': {
        'model': PaymentMethod,
        'filename': 'payment_methods',
        'select_related': ['account'],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Code', 'code'),
            ('Account', 'account.code'),
            ('Requires Reference', 'requires_reference'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'expense_categories': {
        'model': ExpenseCategory,
        'filename': 'expense_categories',
        'select_related': ['parent', 'account'],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Code', 'code'),
            ('Parent', 'parent.name'),
            ('Account', 'account.code'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'currencies': {
        'model': Currency,
        'filename': 'currencies',
        'select_related': [],
        'company_filter': False,
        'columns': [
            ('Code', 'code'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Symbol', 'symbol'),
            ('Exchange Rate', 'exchange_rate'),
            ('Base Currency', 'is_base'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'fiscal_years': {
        'model': FiscalYear,
        'filename': 'fiscal_years',
        'select_related': [],
        'columns': [
            ('Name', 'name'),
            ('Start Date', 'start_date'),
            ('End Date', 'end_date'),
            ('Current', 'is_current'),
            ('Closed', 'is_closed'),
        ],
        'filters': lambda request, qs: qs,
    },
    'leads': {
        'model': Lead,
        'filename': 'leads',
        'select_related': ['assigned_to'],
        'columns': [
            ('Lead #', 'lead_number'),
            ('First Name', 'first_name'),
            ('Last Name', 'last_name'),
            ('Email', 'email'),
            ('Phone', 'phone'),
            ('Company', 'company_name'),
            ('Source', 'get_source_display'),
            ('Status', 'get_status_display'),
            ('Estimated Value', 'estimated_value'),
            ('Assigned To', lambda o: str(o.assigned_to) if o.assigned_to else ''),
            ('Expected Close', 'expected_close_date'),
            ('Probability (%)', 'probability'),
        ],
        'filters': lambda request, qs: _filter_leads(request, qs),
    },
    'campaigns': {
        'model': Campaign,
        'filename': 'campaigns',
        'select_related': ['owner', 'currency'],
        'columns': [
            ('Name', 'name'),
            ('Type', 'get_campaign_type_display'),
            ('Status', 'get_status_display'),
            ('Start Date', 'start_date'),
            ('End Date', 'end_date'),
            ('Budget', 'budget'),
            ('Actual Cost', 'actual_cost'),
            ('Expected Leads', 'expected_leads'),
            ('Actual Leads', 'actual_leads'),
            ('Expected Revenue', 'expected_revenue'),
            ('Actual Revenue', 'actual_revenue'),
            ('Owner', lambda o: str(o.owner) if o.owner else ''),
        ],
        'filters': lambda request, qs: _filter_campaigns(request, qs),
    },
    'tasks': {
        'model': Task,
        'filename': 'tasks',
        'select_related': ['assigned_to', 'related_customer', 'related_lead'],
        'columns': [
            ('Title', 'title'),
            ('Priority', 'get_priority_display'),
            ('Status', 'get_status_display'),
            ('Assigned To', lambda o: str(o.assigned_to) if o.assigned_to else ''),
            ('Due Date', 'due_date'),
            ('Completed Date', 'completed_date'),
            ('Customer', lambda o: str(o.related_customer) if o.related_customer else ''),
            ('Lead', lambda o: str(o.related_lead) if o.related_lead else ''),
        ],
        'filters': lambda request, qs: _filter_tasks(request, qs),
    },
    'employees': {
        'model': Employee,
        'filename': 'employees',
        'select_related': ['department', 'currency'],
        'columns': [
            ('Employee ID', 'employee_id'),
            ('First Name', 'first_name'),
            ('Last Name', 'last_name'),
            ('Email', 'email'),
            ('Phone', 'phone'),
            ('Department', 'department.name'),
            ('Job Title', 'job_title'),
            ('Employment Type', 'get_employment_type_display'),
            ('Status', 'get_employment_status_display'),
            ('Hire Date', 'hire_date'),
            ('Basic Salary', 'basic_salary'),
            ('Currency', 'currency.code'),
        ],
        'filters': lambda request, qs: _filter_employees(request, qs),
    },
    'departments': {
        'model': Department,
        'filename': 'departments',
        'select_related': ['parent', 'manager'],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Code', 'code'),
            ('Parent', 'parent.name'),
            ('Manager', lambda o: str(o.manager) if o.manager else ''),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'attendance': {
        'model': Attendance,
        'filename': 'attendance',
        'select_related': ['employee'],
        'company_field': 'employee__company',
        'columns': [
            ('Employee', lambda o: str(o.employee)),
            ('Date', 'date'),
            ('Check In', 'check_in'),
            ('Check Out', 'check_out'),
            ('Hours Worked', 'hours_worked'),
            ('Overtime', 'overtime_hours'),
            ('Present', 'is_present'),
            ('Late', 'is_late'),
            ('Late Minutes', 'late_minutes'),
        ],
        'filters': lambda request, qs: _filter_attendance(request, qs),
    },
    'leave_requests': {
        'model': LeaveRequest,
        'filename': 'leave_requests',
        'select_related': ['employee', 'leave_type', 'approved_by'],
        'company_field': 'employee__company',
        'columns': [
            ('Employee', lambda o: str(o.employee)),
            ('Leave Type', 'leave_type.name'),
            ('Start Date', 'start_date'),
            ('End Date', 'end_date'),
            ('Days', 'days_requested'),
            ('Status', 'get_status_display'),
            ('Approved By', lambda o: str(o.approved_by) if o.approved_by else ''),
            ('Reason', 'reason'),
        ],
        'filters': lambda request, qs: _filter_leave_requests(request, qs),
    },
    'payroll': {
        'model': Payroll,
        'filename': 'payroll',
        'select_related': ['employee', 'currency'],
        'company_field': 'employee__company',
        'columns': [
            ('Employee', lambda o: str(o.employee)),
            ('Period Start', 'period_start'),
            ('Period End', 'period_end'),
            ('Basic Salary', 'basic_salary'),
            ('Gross Salary', 'gross_salary'),
            ('Total Deductions', 'total_deductions'),
            ('Net Salary', 'net_salary'),
            ('Status', 'get_status_display'),
            ('Payment Date', 'payment_date'),
            ('Currency', 'currency.code'),
        ],
        'filters': lambda request, qs: _filter_payroll(request, qs),
    },
    'suppliers': {
        'model': Supplier,
        'filename': 'suppliers',
        'select_related': ['currency'],
        'columns': [
            ('Supplier ID', 'supplier_id'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Type', 'get_supplier_type_display'),
            ('Email', 'email'),
            ('Phone', 'phone'),
            ('Contact Person', 'contact_person'),
            ('City', 'city'),
            ('Country', 'country'),
            ('Tax Number', 'tax_number'),
            ('Balance', 'current_balance'),
            ('Currency', 'currency.code'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: _filter_suppliers(request, qs),
    },
    'purchase_orders': {
        'model': PurchaseOrder,
        'filename': 'purchase_orders',
        'select_related': ['supplier', 'currency'],
        'columns': [
            ('PO #', 'po_number'),
            ('Supplier', 'supplier.name'),
            ('Date', 'po_date'),
            ('Expected Date', 'expected_date'),
            ('Subtotal', 'subtotal'),
            ('Tax', 'tax_amount'),
            ('Total', 'total_amount'),
            ('Paid', 'paid_amount'),
            ('Balance Due', 'balance_due'),
            ('Status', 'get_status_display'),
            ('Currency', 'currency.code'),
        ],
        'filters': lambda request, qs: _filter_purchase_orders(request, qs),
    },
    'quotes': {
        'model': Quote,
        'filename': 'quotes',
        'select_related': ['customer', 'currency', 'lead'],
        'columns': [
            ('Quote #', 'quote_number'),
            ('Customer', 'customer.name'),
            ('Date', 'quote_date'),
            ('Valid Until', 'valid_until'),
            ('Subtotal', 'subtotal'),
            ('Discount', 'discount_amount'),
            ('Tax', 'tax_amount'),
            ('Total', 'total_amount'),
            ('Status', 'get_status_display'),
            ('Currency', 'currency.code'),
        ],
        'filters': lambda request, qs: _filter_quotes(request, qs),
    },
    'service_categories': {
        'model': ServiceCategory,
        'filename': 'service_categories',
        'select_related': ['parent'],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Code', 'code'),
            ('Parent', 'parent.name'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'services': {
        'model': ServicePricing,
        'filename': 'services',
        'select_related': ['category', 'currency', 'tax_rate'],
        'columns': [
            ('Code', 'code'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Category', 'category.name'),
            ('Pricing Type', 'get_pricing_type_display'),
            ('Base Price', 'base_price'),
            ('Cost Price', 'cost_price'),
            ('Tax Rate', lambda o: str(o.tax_rate.rate) if o.tax_rate else ''),
            ('Currency', 'currency.code'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: _filter_services(request, qs),
    },
    'product_categories': {
        'model': ProductCategory,
        'filename': 'product_categories',
        'select_related': ['parent'],
        'columns': [
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Code', 'code'),
            ('Parent', 'parent.name'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: qs,
    },
    'products': {
        'model': Product,
        'filename': 'products',
        'select_related': ['category', 'currency', 'tax_rate', 'supplier'],
        'columns': [
            ('SKU', 'sku'),
            ('Barcode', 'barcode'),
            ('Name', 'name'),
            ('Name (AR)', 'name_ar'),
            ('Type', 'get_product_type_display'),
            ('Category', 'category.name'),
            ('Selling Price', 'selling_price'),
            ('Cost Price', 'cost_price'),
            ('Quantity on Hand', 'quantity_on_hand'),
            ('Reorder Level', 'reorder_level'),
            ('Supplier', 'supplier.name'),
            ('Currency', 'currency.code'),
            ('Active', 'is_active'),
        ],
        'filters': lambda request, qs: _filter_products(request, qs),
    },
    'journal_entries': {
        'model': JournalEntry,
        'filename': 'journal_entries',
        'select_related': ['fiscal_year'],
        'columns': [
            ('Entry #', 'entry_number'),
            ('Date', 'entry_date'),
            ('Type', 'get_entry_type_display'),
            ('Description', 'description'),
            ('Reference', 'reference'),
            ('Total Debit', 'total_debit'),
            ('Total Credit', 'total_credit'),
            ('Status', 'get_status_display'),
        ],
        'filters': lambda request, qs: _filter_journal_entries(request, qs),
    },
}


# ── Filter functions ────────────────────────────────────────────────────────

def _filter_customers(request, qs):
    if t := request.GET.get('type'):
        qs = qs.filter(customer_type=t)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(name__icontains=s) | Q(email__icontains=s) | Q(customer_id__icontains=s)
        )
    return qs.order_by('name')


def _filter_invoices(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(invoice_number__icontains=s) | Q(customer__name__icontains=s)
        )
    return qs.order_by('-invoice_date')


def _filter_payments(request, qs):
    if t := request.GET.get('type'):
        qs = qs.filter(payment_type=t)
    return qs.order_by('-payment_date')


def _filter_expenses(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    return qs.order_by('-expense_date')


def _filter_accounts(request, qs):
    if t := request.GET.get('type'):
        qs = qs.filter(account_type=t)
    return qs.order_by('code')


def _filter_leads(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    if source := request.GET.get('source'):
        qs = qs.filter(source=source)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(first_name__icontains=s) | Q(last_name__icontains=s) |
            Q(email__icontains=s) | Q(lead_number__icontains=s)
        )
    return qs.order_by('-created_at')


def _filter_campaigns(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    return qs.order_by('-start_date')


def _filter_tasks(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    if priority := request.GET.get('priority'):
        qs = qs.filter(priority=priority)
    return qs.order_by('-created_at')


def _filter_employees(request, qs):
    if dept := request.GET.get('department'):
        qs = qs.filter(department_id=dept)
    if status := request.GET.get('status'):
        qs = qs.filter(employment_status=status)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(first_name__icontains=s) | Q(last_name__icontains=s) |
            Q(email__icontains=s) | Q(employee_id__icontains=s)
        )
    return qs.order_by('first_name')


def _filter_attendance(request, qs):
    from datetime import date as _date
    today = _date.today()
    start = request.GET.get('start_date', today.replace(day=1).isoformat())
    end = request.GET.get('end_date', today.isoformat())
    qs = qs.filter(date__gte=start, date__lte=end)
    return qs.order_by('-date', 'employee__first_name')


def _filter_leave_requests(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    return qs.order_by('-created_at')


def _filter_payroll(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(employee__first_name__icontains=s) |
            Q(employee__last_name__icontains=s) |
            Q(employee__employee_id__icontains=s)
        )
    return qs.order_by('-period_start')


def _filter_suppliers(request, qs):
    if t := request.GET.get('type'):
        qs = qs.filter(supplier_type=t)
    if s := request.GET.get('search'):
        qs = qs.filter(
            Q(name__icontains=s) | Q(email__icontains=s) | Q(supplier_id__icontains=s)
        )
    return qs.order_by('name')


def _filter_purchase_orders(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    return qs.order_by('-po_date')


def _filter_quotes(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    return qs.order_by('-quote_date')


def _filter_services(request, qs):
    if s := request.GET.get('search', '').strip():
        qs = qs.filter(
            Q(name__icontains=s) | Q(code__icontains=s) | Q(name_ar__icontains=s)
        )
    if cat := request.GET.get('category'):
        qs = qs.filter(category_id=cat)
    return qs.order_by('name')


def _filter_products(request, qs):
    if s := request.GET.get('search', '').strip():
        qs = qs.filter(
            Q(name__icontains=s) | Q(sku__icontains=s) |
            Q(barcode__icontains=s) | Q(name_ar__icontains=s)
        )
    if cat := request.GET.get('category'):
        qs = qs.filter(category_id=cat)
    if ptype := request.GET.get('type'):
        qs = qs.filter(product_type=ptype)
    return qs.order_by('name')


def _filter_journal_entries(request, qs):
    if status := request.GET.get('status'):
        qs = qs.filter(status=status)
    if entry_type := request.GET.get('type'):
        qs = qs.filter(entry_type=entry_type)
    if df := request.GET.get('date_from'):
        qs = qs.filter(entry_date__gte=df)
    if dt := request.GET.get('date_to'):
        qs = qs.filter(entry_date__lte=dt)
    return qs.order_by('-entry_date', '-entry_number')


# ── Export functions ────────────────────────────────────────────────────────

MAX_EXPORT_ROWS = 50_000
EXPORT_CACHE_TTL = 300  # 5 minutes


def _cache_key(entity, company_id, fmt, query_params):
    """Build a deterministic cache key for an export request."""
    # Sort params for consistency regardless of order
    params = sorted(
        (k, v) for k, v in query_params.items() if k != 'format'
    )
    raw = f'export:{entity}:{company_id}:{fmt}:{params}'
    digest = hashlib.sha256(raw.encode()).hexdigest()[:16]
    return f'export:{entity}:{company_id}:{digest}'


def invalidate_export_cache(entity, company_id):
    """Delete all cached exports for an entity+company.

    Called from model signals when data changes.
    Uses a generation counter so we don't need to track every key.
    """
    gen_key = f'export_gen:{entity}:{company_id}'
    try:
        cache.incr(gen_key)
    except ValueError:
        # Key doesn't exist yet — initialise it
        cache.set(gen_key, 1, None)


def _versioned_cache_key(entity, company_id, fmt, query_params):
    """Cache key incorporating a generation counter for invalidation."""
    gen_key = f'export_gen:{entity}:{company_id}'
    gen = cache.get(gen_key, 0)
    base = _cache_key(entity, company_id, fmt, query_params)
    return f'{base}:g{gen}'


def get_export_queryset(entity, request, company):
    """Build the filtered queryset for the given entity."""
    config = EXPORT_REGISTRY.get(entity)
    if not config:
        return None, None

    model = config['model']
    company_filter = config.get('company_filter', True)

    if company_filter is False:
        qs = model.objects.all()
    else:
        company_field = config.get('company_field', 'company')
        qs = model.objects.filter(**{company_field: company})

    sr = config.get('select_related', [])
    if sr:
        qs = qs.select_related(*sr)

    filter_fn = config.get('filters')
    if filter_fn:
        qs = filter_fn(request, qs)

    return qs[:MAX_EXPORT_ROWS], config


def _build_csv_bytes(queryset, config):
    """Generate CSV content as bytes (cacheable)."""
    buf = BytesIO()
    import io
    wrapper = io.TextIOWrapper(buf, encoding='utf-8', newline='')
    wrapper.write('\ufeff')  # BOM
    columns = config['columns']
    writer = csv.writer(wrapper)
    writer.writerow([col[0] for col in columns])
    for obj in queryset.iterator():
        writer.writerow([_resolve_field(obj, col[1]) for col in columns])
    wrapper.flush()
    wrapper.detach()  # release buf without closing it
    return buf.getvalue()


def _build_excel_bytes(queryset, config):
    """Generate Excel content as bytes (cacheable)."""
    wb = Workbook()
    ws = wb.active
    ws.title = config['filename'][:31]

    columns = config['columns']
    headers = [col[0] for col in columns]

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2E86AB', end_color='2E86AB', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(bottom=Side(style='thin', color='CCCCCC'))

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, obj in enumerate(queryset.iterator(), 2):
        for col_idx, col_def in enumerate(columns, 1):
            value = _resolve_field(obj, col_def[1])
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border

    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 3, 50)

    ws.freeze_panes = 'A2'

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(queryset, config, filename, *, cache_key=None):
    """Export queryset to CSV with UTF-8 BOM for Excel compatibility."""
    if cache_key:
        data = cache.get(cache_key)
    else:
        data = None

    if data is None:
        data = _build_csv_bytes(queryset, config)
        if cache_key:
            cache.set(cache_key, data, EXPORT_CACHE_TTL)

    response = HttpResponse(data, content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def export_excel(queryset, config, filename, *, cache_key=None):
    """Export queryset to Excel (.xlsx) with styled headers."""
    if cache_key:
        data = cache.get(cache_key)
    else:
        data = None

    if data is None:
        data = _build_excel_bytes(queryset, config)
        if cache_key:
            cache.set(cache_key, data, EXPORT_CACHE_TTL)

    response = HttpResponse(
        data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response
